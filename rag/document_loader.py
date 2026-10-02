import os
import io
import tempfile
from typing import List, Dict, Any, Union
import fitz  # PyMuPDF
import docx
import pandas as pd

class DocumentChunk:
    def __init__(self, text: str, metadata: Dict[str, Any]):
        self.text = text
        self.metadata = metadata

    def to_dict(self) -> Dict[str, Any]:
        return {"text": self.text, "metadata": self.metadata}

class DocumentLoader:
    """Universal Document Loader supporting PDF, DOCX, TXT, MD, CSV files and Streamlit UploadedFiles."""

    @staticmethod
    def load_pdf(file_bytes: bytes, filename: str) -> List[DocumentChunk]:
        chunks = []
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        total_pages = len(doc)
        for page_idx in range(total_pages):
            page = doc.load_page(page_idx)
            text = page.get_text("text").strip()
            if text:
                chunks.append(
                    DocumentChunk(
                        text=text,
                        metadata={
                            "source": filename,
                            "page": page_idx + 1,
                            "total_pages": total_pages,
                            "file_type": "pdf"
                        }
                    )
                )
        return chunks

    @staticmethod
    def load_docx(file_bytes: bytes, filename: str) -> List[DocumentChunk]:
        chunks = []
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        table_texts = []
        for table_idx, table in enumerate(doc.tables):
            rows_data = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                if any(row_cells):
                    rows_data.append(" | ".join(row_cells))
            if rows_data:
                table_texts.append(f"[Table {table_idx+1}]\n" + "\n".join(rows_data))
        
        all_text = "\n\n".join(paragraphs + table_texts)
        if all_text:
            chunks.append(
                DocumentChunk(
                    text=all_text,
                    metadata={
                        "source": filename,
                        "page": 1,
                        "total_pages": 1,
                        "file_type": "docx"
                    }
                )
            )
        return chunks

    @staticmethod
    def load_txt(file_bytes: bytes, filename: str) -> List[DocumentChunk]:
        text = ""
        for encoding in ["utf-8", "latin-1", "cp1252"]:
            try:
                text = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        text = text.strip()
        if not text:
            return []
        return [
            DocumentChunk(
                text=text,
                metadata={
                    "source": filename,
                    "page": 1,
                    "total_pages": 1,
                    "file_type": "txt"
                }
            )
        ]

    @staticmethod
    def load_csv(file_bytes: bytes, filename: str) -> List[DocumentChunk]:
        try:
            df = pd.read_csv(io.BytesIO(file_bytes))
            preview_str = df.head(10).to_string(index=False)
            summary_stats = f"CSV Dataset: {filename}\nColumns: {', '.join(df.columns.tolist())}\nTotal Rows: {len(df)}\n\nPreview (First 10 rows):\n{preview_str}"
            
            if not df.select_dtypes(include='number').empty:
                summary_stats += "\n\nDescriptive Statistics:\n" + df.describe().to_string()

            return [
                DocumentChunk(
                    text=summary_stats,
                    metadata={
                        "source": filename,
                        "page": 1,
                        "total_pages": 1,
                        "file_type": "csv",
                        "row_count": len(df)
                    }
                )
            ]
        except Exception:
            return DocumentLoader.load_txt(file_bytes, filename)

    @classmethod
    def load_file(cls, file_source: Union[str, Any], filename: str = None) -> List[DocumentChunk]:
        """Loads file from filepath or Streamlit UploadedFile object."""
        if hasattr(file_source, "read") and hasattr(file_source, "name"):
            filename = filename or file_source.name
            file_bytes = file_source.getvalue() if hasattr(file_source, "getvalue") else file_source.read()
        elif isinstance(file_source, str) and os.path.exists(file_source):
            filename = filename or os.path.basename(file_source)
            with open(file_source, "rb") as f:
                file_bytes = f.read()
        elif isinstance(file_source, bytes):
            filename = filename or "uploaded_document"
            file_bytes = file_source
        else:
            raise ValueError(f"Unsupported file source type: {type(file_source)}")

        ext = os.path.splitext(filename)[1].lower()
        if ext == ".pdf":
            return cls.load_pdf(file_bytes, filename)
        elif ext in [".docx", ".doc"]:
            return cls.load_docx(file_bytes, filename)
        elif ext in [".txt", ".md", ".json", ".log"]:
            return cls.load_txt(file_bytes, filename)
        elif ext in [".csv", ".tsv"]:
            return cls.load_csv(file_bytes, filename)
        else:
            return cls.load_txt(file_bytes, filename)
